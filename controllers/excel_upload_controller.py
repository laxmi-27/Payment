import config
import ast
import gc
import json
import math
import re
import os
import time
import openpyxl
from flask_login import current_user, login_required
from flask import Blueprint, request, render_template, redirect, url_for, jsonify

from model.Log import LogHelper
from model.FolderAndFile import FolderAndFile
from model.HoldReleaseSummary import infer_target, FIXED_HOLDS
from datetime import datetime

excel_bp = Blueprint('excel', __name__)

_DATE_FMTS = ('%d-%m-%Y', '%d/%m/%Y', '%Y-%m-%d', '%m/%d/%Y', '%d.%m.%Y')


def _norm_hold_name(s):
    return re.sub(r'[\s_]+', '', str(s or '')).lower()


def _is_hold_release_note(details_lower):
    """
    True for notes like 'Hold Release Note', 'Testing deposit Release Note',
    'Maintenance Release Note', 'SD Release Note' ...
    Any '... release ...' note that is NOT a GST release.
    """
    return 'release' in details_lower and 'gst' not in details_lower


def _hold_target_from_text(cursor, details):
    """
    Read the hold column from the note text. Spaces / underscores / case are ignored:
      'Testing_ Deposit_RR_release_note' -> Testing_ Deposit_RR   (hold type)
      'Maintaninece_RR_release_note'     -> Maintaninece_RR       (hold type)
      'SD_RR_release_note'               -> SD_RR                 (hold type)
      'SD_Amount_release_note'           -> SD_Amount             (fixed column)
      'Hydro_Testing_release_note'       -> Hydro_Testing         (fixed column)
      'Hydro_Testing_RR_release_note'    -> Hydro_Testing         (fixed column)
    Returns (hold_type_id, fixed_hold) or (None, None).
    """
    key = _norm_hold_name(details)
    if not key:
        return None, None
    for fx in FIXED_HOLDS:                                   # sdamount / hydrotesting
        if _norm_hold_name(fx) in key:
            return None, fx
    if 'hydro' in key:
        return None, 'Hydro_Testing'
    cursor.execute("SELECT hold_type_id, hold_type FROM hold_types")
    types = cursor.fetchall()
    best = None                                              # longest matching name wins
    for tid, tname in types:
        n = _norm_hold_name(tname)
        n_short = n[:-2] if n.endswith('rr') else n          # 'testingdepositrr' -> 'testingdeposit'
        hit = len(n) if (n and n in key) else (len(n_short) if (len(n_short) > 3 and n_short in key) else 0)
        if hit and (best is None or hit > best[0]):
            best = (hit, tid)
    if best:
        return best[1], None
    if 'maint' in key:                                       # maintenance / maintainance / maintaninece
        for tid, tname in types:
            if 'maint' in _norm_hold_name(tname):
                return tid, None
    if re.match(r'^(sd|securitydeposit)(release|amt)', key):  # plain 'SD Release Note'
        return None, 'SD_Amount'
    return None, None


def _resolve_hold_release_target(cursor, entry, contractor_id, pmc_no, invoice_no, amount, details=''):
    """
    Decide which hold column a release note belongs to:
      1. Excel column 'Release_Column' (e.g. Testing_ Deposit_RR)
      2. The note text itself ('Testing deposit Release Note')
      3. Inference from invoice no + amount (HoldReleaseSummary.infer_target)
    Returns (hold_type_id, fixed_hold).
    """
    # Optional extra Excel column. Header can be written 'Release_Column', 'Release Column',
    # 'Released Column' ...  (do NOT use the word 'hold' in the header - the parser would
    # treat it as a hold amount column).
    name = ''
    for _k, _v in entry.items():
        if _norm_hold_name(_k) in ('releasecolumn', 'releasedcolumn', 'releasecol', 'releasefrom') and _v:
            name = _v
            break
    if str(name).strip():
        key = _norm_hold_name(name)
        for fx in FIXED_HOLDS:                       # SD_Amount / Hydro_Testing
            if _norm_hold_name(fx) == key:
                return None, fx
        t, f = _hold_target_from_text(cursor, name)  # Testing_ Deposit_RR / SD_RR / Maintaninece_RR
        if t or f:
            return t, f
    t, f = _hold_target_from_text(cursor, details)
    if t or f:
        return t, f
    return infer_target(cursor, contractor_id, pmc_no, invoice_no, amount)


def _to_float(v):
    """Parse Indian-format amounts like 1,23,456.78 or plain numbers. Returns 0.0 on failure."""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(',', '').replace('₹', '').replace('Rs.', '').replace('Rs', '').strip()
    try:
        return float(s)
    except (ValueError, TypeError):
        return 0.0


def _parse_date(v):
    """Parse date from datetime object, Excel serial, or multiple string formats. Returns YYYY-MM-DD or None."""
    if v is None:
        return None
    if hasattr(v, 'strftime'):
        return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    for fmt in _DATE_FMTS:
        try:
            return datetime.strptime(s, fmt).strftime('%Y-%m-%d')
        except ValueError:
            continue
    return None


def _is_total_row(entry):
    """Return True if row looks like a total/summary/blank row that should be skipped."""
    pmc = str(entry.get('PMC_No') or '').strip().lower()
    if pmc in ('', 'total', 'grand total', 'subtotal', 'sub total', 'net total'):
        return True
    if any(kw in pmc for kw in ('total', 'grand', 'subtotal')):
        return True
    vals = [v for k, v in entry.items() if k != 'Row Number' and v is not None and str(v).strip() != '']
    return len(vals) < 3


def _clean_invoice_no(v):
    """Return None for blank/dash/zero invoice numbers, else the stripped string."""
    if v is None:
        return None
    s = str(v).strip()
    if s in ('-', '--', '—', 'N/A', 'n/a', 'NA', 'nil', 'Nil', 'NIL', '0', ''):
        return None
    return s


def _clean_utr(v):
    """Return None for blank UTR values so the UNIQUE constraint handles dedup."""
    if v is None:
        return None
    s = str(v).strip()
    return s if s and s not in ('-', 'N/A', 'NA', 'nil', 'NIL', '0') else None


# ── Column header normalization ────────────────────────────────────

# Maps lowercased-no-space column names → canonical field names used in save_data
_HEADER_NORM = {
    'invoicedate': 'Invoice_Date',
    'date': 'Invoice_Date',
    'billdate': 'Invoice_Date',
    'invoiceno': 'Invoice_No',
    'invoicenumber': 'Invoice_No',
    'billno': 'Invoice_No',
    'billnumber': 'Invoice_No',
    'invoicedetails': 'Invoice_Details',
    'invoicedetail': 'Invoice_Details',
    'description': 'Invoice_Details',
    'details': 'Invoice_Details',
    'workdescription': 'Invoice_Details',
    'particulars': 'Invoice_Details',
    'narration': 'Invoice_Details',
    'basicamount': 'Basic_Amount',
    'basicamt': 'Basic_Amount',
    'baseamount': 'Basic_Amount',
    'amount': 'Amount',
    'debitamount': 'Debit_Amount',
    'debitamt': 'Debit_Amount',
    'debit': 'Debit_Amount',
    'afterdebitamount': 'After_Debit_Amount',
    'afterdebit': 'After_Debit_Amount',
    'gstamount': 'GST_Amount',
    'gstamt': 'GST_Amount',
    'gst': 'GST_Amount',
    'tdsamount': 'TDS_Amount',
    'tdsamt': 'TDS_Amount',
    'tds': 'TDS_Amount',
    'finalamount': 'Final_Amount',
    'finalamt': 'Final_Amount',
    'netamount': 'Final_Amount',
    'netamt': 'Final_Amount',
    'payableamount': 'Final_Amount',
    'payableamt': 'Final_Amount',
    'totalamount': 'Total_Amount',
    'totalamt': 'Total_Amount',
    'paidamount': 'Total_Amount',
    'paidamt': 'Total_Amount',
    'amountpaid': 'Total_Amount',
    'paymentamount': 'Payment_Amount',
    'paymentamt': 'Payment_Amount',
    'tdspaymentamount': 'TDS_Payment_Amount',
    'tdspaymentamt': 'TDS_Payment_Amount',
    'utrno': 'UTR',
    'utr': 'UTR',
    'referenceno': 'UTR',
    'refno': 'UTR',
    'transactionid': 'UTR',
    'transid': 'UTR',
    'pmcno': 'PMC_No',
    'pmcnumber': 'PMC_No',
    'pmc': 'PMC_No',
    'villagename': 'Village_Name',
    'village': 'Village_Name',
    'sdamount': 'SD_Amount',
    'sdamt': 'SD_Amount',
    'hydrotesting': 'Hydro_Testing',
    'hydro': 'Hydro_Testing',
    'gstsdamount': 'GST_SD_Amount',
    'holdamount': 'Hold_Amount',
    'holdamt': 'Hold_Amount',
    'contractorname': 'Contractor_Name',
    'subcontractorname': 'Contractor_Name',
    'contractor': 'Contractor_Name',
    'villagewiseadvance': 'Village_Wise_Advance',
    'villageadvance': 'Village_Wise_Advance',
    'villagewisedvance': 'Village_Wise_Advance',
    'villagwiseadvance': 'Village_Wise_Advance',
}

# Maps lowercased-no-space metadata labels → file_info key
_META_LABELS = {
    'subcontractor': 'Subcontractor',
    'subcontractorname': 'Subcontractor',
    'subcontractors': 'Subcontractor',      # plural form
    'subcontractorsname': 'Subcontractor',  # plural + name
    'subcontactor': 'Subcontractor',        # common typo (missing r)
    'subcontactorname': 'Subcontractor',    # common typo + name
    'subcontracter': 'Subcontractor',       # alternate misspelling
    'subcontractername': 'Subcontractor',   # alternate misspelling + name
    'contractorname': 'Subcontractor',
    'contractor': 'Subcontractor',
    'name': 'Subcontractor',
    'state': 'State',
    'statename': 'State',
    'district': 'District',
    'districtname': 'District',
    'block': 'Block',
    'blockname': 'Block',
    'village': 'Village',
    'villagename': 'Village',
}

# Ordered longest-first for prefix matching (avoids "state" matching before "statename")
_META_LABEL_PREFIXES = sorted(_META_LABELS.keys(), key=len, reverse=True)


def _norm_key(v):
    """Lowercase + remove spaces/underscores/slashes/colons for fuzzy matching."""
    return re.sub(r'[\s_/\-:]+', '', str(v).lower().strip())


def _map_header(raw):
    """Translate a raw Excel column header to a canonical field name."""
    return _HEADER_NORM.get(_norm_key(raw), str(raw).strip())


def _detect_header_row(sheet):
    """
    Scan rows 1-25 and return the row index that best matches known column headers.
    Returns (header_row_idx, data_start_row_idx).
    Falls back to row 5 / row 6 if nothing matches.
    """
    best_row, best_score = 5, 0
    for r in range(1, min(26, sheet.max_row + 1)):
        score = 0
        for c in range(1, sheet.max_column + 1):
            v = sheet.cell(row=r, column=c).value
            if v and _norm_key(str(v)) in _HEADER_NORM:
                score += 1
        if score > best_score:
            best_score, best_row = score, r
    return best_row, best_row + 1


def _detect_metadata(sheet, header_row):
    """
    Scan rows before the header row for label: value pairs (metadata).
    Handles two patterns:
      1. Label in one cell, value in a nearby column (searches up to 5 cols right)
      2. Label + value combined in the same cell (e.g. "Subcontractors Raghuveer Singh & Sons")
    Returns a dict with keys from _META_LABELS values.
    """
    info = {}
    max_col = min(sheet.max_column, 12)

    for r in range(1, header_row):
        for c in range(1, max_col + 1):
            v = sheet.cell(row=r, column=c).value
            if v is None:
                continue
            cell_str = str(v).strip()
            nk = _norm_key(cell_str)

            # ── Pattern 1: exact label → scan right up to 5 cols for value ──
            if nk in _META_LABELS:
                meta_key = _META_LABELS[nk]
                if meta_key not in info:
                    val = None
                    for vc in range(c + 1, min(c + 6, sheet.max_column + 1)):
                        candidate = sheet.cell(row=r, column=vc).value
                        if candidate is not None and str(candidate).strip():
                            val = candidate
                            break
                    if val is None and c > 1:
                        val = sheet.cell(row=r, column=c - 1).value
                    if val is not None and str(val).strip():
                        info[meta_key] = str(val).strip()
                continue

            # ── Pattern 2: label+value in same cell (e.g. "Subcontractors Foo Bar") ──
            for lk in _META_LABEL_PREFIXES:
                if nk.startswith(lk) and len(nk) > len(lk):
                    meta_key = _META_LABELS[lk]
                    if meta_key not in info:
                        # Reconstruct value by stripping the label prefix from original text
                        low = cell_str.lower()
                        # Find the raw label length by counting non-space chars
                        raw_skip = 0
                        matched = 0
                        for ch in cell_str:
                            if matched >= len(lk):
                                break
                            raw_skip += 1
                            if ch not in (' ', '\t', '_', '-', '/'):
                                matched += 1
                        remainder = cell_str[raw_skip:].lstrip(' :-\t').strip()
                        if remainder:
                            info[meta_key] = remainder
                    break

    return info


# ── Fuzzy location lookup helpers ────────────────────────────────
# Strips trailing noise words that appear in Excel metadata but not in DB names.
# e.g. "Muzaffarnagar UP" → "Muzaffarnagar", "Uttar Pradesh UP" → "Uttar Pradesh"
# Village-word variants: handles common misspellings and punctuation-attached forms
# e.g. "Village-" (dash glued), "Vllage", "Villgae", "Vilage"
_VILLAGE_TOKENS = {'village', 'vllage', 'vilage', 'villge', 'villag', 'vialge', 'villgae', 'gram', 'gaon'}

_NOISE_SUFFIXES = (
    ' up', ' mp', ' hp', ' gj', ' rj', ' mh', ' wb', ' ap', ' ts', ' ka',
    ' tn', ' kl', ' od', ' jk', ' br', ' uk', ' cg', ' jh', ' as', ' ga',
    ' uttar pradesh', ' madhya pradesh', ' rajasthan', ' maharashtra',
    ' gujarat', ' bihar', ' jharkhand', ' uttarakhand', ' haryana',
    ' himachal pradesh', ' punjab', ' delhi', ' odisha', ' west bengal',
)

def _fuzzy_clean(name):
    """Return a cleaned version of a location name for fuzzy matching."""
    s = (name or '').strip().lower()
    for suffix in _NOISE_SUFFIXES:
        if s.endswith(suffix):
            s = s[: -len(suffix)].strip()
    return s.title()


def _fuzzy_lookup_state(cursor, name):
    """Try exact SP, then fuzzy fallback via LIKE on cleaned name."""
    cursor.callproc('CheckStateExists', [name])
    row = None
    for r in cursor.stored_results():
        row = r.fetchone()
    if row:
        return row
    cleaned = _fuzzy_clean(name)
    if cleaned and cleaned.lower() != name.strip().lower():
        cursor.execute(
            "SELECT * FROM states WHERE LOWER(State_Name) LIKE %s LIMIT 1",
            ('%' + cleaned.lower() + '%',)
        )
        row = cursor.fetchone()
    return row


def _fuzzy_lookup_district(cursor, name, state_id):
    """Try exact SP, then fuzzy fallback via LIKE on cleaned name."""
    cursor.callproc('GetDistrictByNameAndState', [name, state_id])
    row = None
    for r in cursor.stored_results():
        row = r.fetchone()
    if row:
        return row
    cleaned = _fuzzy_clean(name)
    cursor.execute(
        "SELECT * FROM districts WHERE State_Id = %s AND LOWER(District_Name) LIKE %s LIMIT 1",
        (state_id, '%' + cleaned.lower() + '%')
    )
    return cursor.fetchone()


def _fuzzy_lookup_block(cursor, name, district_id, connection=None):
    """
    Try exact SP, then fuzzy LIKE fallback.
    If still not found AND connection is provided, auto-create the block
    and return (row, True) to signal a new block was added.
    Returns (row, False) when found, (None, False) when not found and no connection.
    """
    cursor.callproc('GetBlockByNameAndDistrict', [name, district_id])
    row = None
    for r in cursor.stored_results():
        row = r.fetchone()
    if row:
        return row, False
    cleaned = _fuzzy_clean(name)
    cursor.execute(
        "SELECT * FROM blocks WHERE District_Id = %s AND LOWER(Block_Name) LIKE %s LIMIT 1",
        (district_id, '%' + cleaned.lower() + '%')
    )
    row = cursor.fetchone()
    if row:
        return row, False
    # Auto-create the block if a live connection was provided
    if connection and name and name.strip():
        block_name = name.strip()
        cursor.callproc('SaveBlock', [block_name, district_id])
        connection.commit()
        cursor.execute(
            "SELECT * FROM blocks WHERE Block_Name = %s AND District_Id = %s LIMIT 1",
            (block_name, district_id)
        )
        row = cursor.fetchone()
        return row, True   # True = newly created
    return None, False


# ── uploaded_files DB helpers ─────────────────────────────────────

def _get_all_uploads():
    conn = cursor = None
    try:
        conn   = config.get_db_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT * FROM uploaded_files ORDER BY uploaded_date DESC")
        return cursor.fetchall()
    except Exception as e:
        print("_get_all_uploads error:", e)
        return []
    finally:
        if cursor: cursor.close()
        if conn:   conn.close()


def _find_by_stored_name(stored_name):
    conn = cursor = None
    try:
        conn   = config.get_db_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            "SELECT * FROM uploaded_files WHERE stored_name = %s LIMIT 1",
            (stored_name,)
        )
        return cursor.fetchone()
    except Exception as e:
        print("_find_by_stored_name error:", e)
        return None
    finally:
        if cursor: cursor.close()
        if conn:   conn.close()


def _insert_upload(original_name, stored_name, total_rows=0):
    """Insert a new upload record and return its file_id (or None on failure)."""
    conn = cursor = None
    try:
        conn   = config.get_db_connection()
        cursor = conn.cursor()
        cursor.execute(
            "INSERT INTO uploaded_files (original_name, stored_name, total_rows) VALUES (%s, %s, %s)",
            (original_name, stored_name, total_rows)
        )
        conn.commit()
        return cursor.lastrowid
    except Exception as e:
        print("_insert_upload error:", e)
        return None
    finally:
        if cursor: cursor.close()
        if conn:   conn.close()


def _auto_rename(original_name):
    """Find the next available versioned name, checking both disk and DB."""
    folder = FolderAndFile.get_upload_folder()
    name, ext = os.path.splitext(original_name)
    counter = 2
    while True:
        new_name = f"{name}_{counter}{ext}"
        on_disk = os.path.exists(os.path.join(folder, new_name))
        in_db   = bool(_find_by_stored_name(new_name))
        if not on_disk and not in_db:
            return new_name
        counter += 1


def _count_excel_rows(filepath):
    wb = None
    try:
        wb    = openpyxl.load_workbook(filepath, data_only=True)
        sheet = wb.active
        count = 0
        for i in range(6, sheet.max_row + 1):
            if sheet.cell(row=i, column=1).value is not None:
                row_vals = [sheet.cell(row=i, column=j).value
                            for j in range(1, sheet.max_column + 1)]
                if sum(1 for v in row_vals if v) >= 4:
                    count += 1
        return count
    except Exception:
        return 0
    finally:
        if wb is not None:
            wb.close()
        del wb
        gc.collect()


def _safe_remove(filepath, retries=3, delay=0.6):
    """Delete a file with retry on Windows file-lock (WinError 32)."""
    last_err = None
    for attempt in range(retries):
        try:
            if os.path.exists(filepath):
                os.remove(filepath)
            return True
        except PermissionError as e:
            last_err = e
            if attempt < retries - 1:
                gc.collect()
                time.sleep(delay)
        except OSError as e:
            raise
    raise PermissionError(f"File is locked after {retries} attempts: {last_err}")


def _safe_save_upload(file_storage, filepath, retries=3, delay=0.6):
    """Save a Werkzeug FileStorage to filepath, retrying on PermissionError."""
    last_err = None
    for attempt in range(retries):
        try:
            file_storage.stream.seek(0)
            file_storage.save(filepath)
            return True
        except PermissionError as e:
            last_err = e
            if attempt < retries - 1:
                gc.collect()
                time.sleep(delay)
        except Exception:
            raise
    raise PermissionError(f"Cannot write file after {retries} attempts: {last_err}")


def _parse_excel_file(filepath):
    """
    Open an Excel workbook, detect the header row, extract metadata and all
    data rows.  Always closes the workbook regardless of success or failure.

    Returns a dict:
        file_info       – {Subcontractor, State, District, Block, Village}
        variables       – {canonical_field_name: column_index}   (1-based)
        hold_columns_raw– [{column_name, column_number}]  (no hold_type_id)
        data            – list of row-dicts keyed by canonical field names
    """
    wb = None
    try:
        wb    = openpyxl.load_workbook(filepath, data_only=True)
        sheet = wb.active

        header_row, data_start = _detect_header_row(sheet)

        file_info = _detect_metadata(sheet, header_row)
        file_info.setdefault('Subcontractor', None)
        file_info.setdefault('State',         None)
        file_info.setdefault('District',      None)
        file_info.setdefault('Block',         None)

        # ── Map column headers → canonical names ─────────────────────────
        variables        = {}   # canonical → column-index
        hold_columns_raw = []   # hold cols without DB-resolved hold_type_id

        for j in range(1, sheet.max_column + 1):
            raw_val = sheet.cell(row=header_row, column=j).value
            if not raw_val:
                continue
            raw       = str(raw_val).strip()
            canonical = _map_header(raw)
            variables[canonical] = j
            if 'hold' in raw.lower():
                hold_columns_raw.append({'column_name': canonical, 'column_number': j})

        # ── Read data rows ────────────────────────────────────────────────
        rows = []
        last_pmc = None   # fill-down: carry PMC forward when cell is blank
        pmc_col = variables.get('PMC_No')
        for i in range(data_start, sheet.max_row + 1):
            raw_pmc = sheet.cell(row=i, column=1).value if pmc_col else None
            pmc_val = str(raw_pmc).strip() if raw_pmc is not None else ''

            # If PMC cell is empty, only fill-down for genuine continuation rows
            if not pmc_val:
                if last_pmc is None:
                    continue
                # Skip rows whose text contains summary/total keywords
                _summary_kws = ('total', 'grand total', 'balance', 'subtotal', 'net total', 'paid')
                row_text = ' '.join(
                    str(sheet.cell(row=i, column=c).value or '').lower()
                    for c in range(1, min(sheet.max_column + 1, 25))
                )
                if any(kw in row_text for kw in _summary_kws):
                    continue
                # A fill-down row is only valid if it has a UTR (continuation payment)
                # OR explicit invoice data (Invoice_Details / Invoice_No / Invoice_Date).
                # Rows with only financial amounts and no UTR/details are totals/summaries.
                utr_col      = variables.get('UTR')
                det_col      = variables.get('Invoice_Details')
                inv_no_col   = variables.get('Invoice_No')
                inv_date_col = variables.get('Invoice_Date')
                has_utr = utr_col and sheet.cell(row=i, column=utr_col).value not in (None, '', 0)
                has_inv = (
                    (det_col      and sheet.cell(row=i, column=det_col).value      not in (None, '', 0)) or
                    (inv_no_col   and sheet.cell(row=i, column=inv_no_col).value   not in (None, '', 0)) or
                    (inv_date_col and sheet.cell(row=i, column=inv_date_col).value not in (None, '', 0))
                )
                if not has_utr and not has_inv:
                    continue
                # Fill-down PMC from the last seen value
                pmc_val = last_pmc
            else:
                last_pmc = pmc_val

            row_data = {'Row Number': i}
            for var_name, col_num in variables.items():
                cell_val = sheet.cell(row=i, column=col_num).value
                if isinstance(cell_val, float) and math.isnan(cell_val):
                    cell_val = ''
                elif isinstance(cell_val, str):
                    cell_val = cell_val.strip()
                row_data[var_name] = cell_val
            # Override PMC_No with fill-down value when cell was blank
            if not row_data.get('PMC_No'):
                row_data['PMC_No'] = pmc_val
            if sum(1 for v in row_data.values() if v) >= 4:
                rows.append(row_data)

        return {
            'file_info':        file_info,
            'variables':        variables,
            'hold_columns_raw': hold_columns_raw,
            'data':             rows,
        }
    finally:
        if wb is not None:
            wb.close()
        del wb
        gc.collect()


# ── Upload ────────────────────────────────────────────────────────

@excel_bp.route('/upload_excel_file', methods=['GET', 'POST'])
@login_required
def upload():
    if request.method == 'POST':
        is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'
        file    = request.files.get('file')
        mode    = request.form.get('mode', 'new')   # 'new' | 'replace' | 'new_version'

        if not file or not file.filename.endswith(('.xlsx', '.xls')):
            if is_ajax:
                return jsonify({"ok": False, "error": "Please select a valid Excel file (.xlsx or .xls)."}), 400
            return render_template('uploadExcelFile.html')

        original_name = file.filename
        folder        = FolderAndFile.get_upload_folder()

        # ── Duplicate detection (server-side, DB is source of truth) ───────
        if mode == 'new':
            in_db = _find_by_stored_name(original_name)
            if in_db:
                # Tell the client a duplicate was found WITHOUT saving the file
                if is_ajax:
                    return jsonify({"ok": False, "duplicate": True, "filename": original_name})
                return render_template('uploadExcelFile.html')

        # ── Determine stored name ────────────────────────────────────────
        if mode == 'new_version':
            stored_name = _auto_rename(original_name)
        else:
            stored_name = original_name          # 'replace' keeps the same name

        # ── Save file to disk ────────────────────────────────────────────
        filepath   = FolderAndFile.get_upload_path(stored_name)
        _safe_save_upload(file, filepath)
        try:
            file.stream.close()
        except Exception:
            pass
        total_rows = _count_excel_rows(filepath)

        # ── Update DB record and get file_id ─────────────────────────────
        active_file_id = None
        if mode == 'replace':
            existing = _find_by_stored_name(original_name)
            if existing:
                active_file_id = existing['file_id']
                conn = cursor = None
                try:
                    conn   = config.get_db_connection()
                    cursor = conn.cursor()
                    cursor.execute(
                        "UPDATE uploaded_files "
                        "SET updated_date = NOW(), total_rows = %s "
                        "WHERE file_id = %s",
                        (total_rows, active_file_id)
                    )
                    conn.commit()
                except Exception as e:
                    print("replace update error:", e)
                finally:
                    if cursor: cursor.close()
                    if conn:   conn.close()
            else:
                active_file_id = _insert_upload(original_name, stored_name, total_rows)
        else:
            active_file_id = _insert_upload(original_name, stored_name, total_rows)

        LogHelper.log_action(
            "Upload Excel File",
            f"User {current_user.id} uploaded '{stored_name}' (mode={mode})"
        )

        redirect_url = url_for('excel.show_table', filename=stored_name,
                                file_id=active_file_id or '')
        if is_ajax:
            return jsonify({"ok": True, "redirect": redirect_url})
        return redirect(redirect_url)

    return render_template('uploadExcelFile.html')


# ── API: check duplicate filename ────────────────────────────────

@excel_bp.route('/api/check_excel_name', methods=['GET'])
@login_required
def check_excel_name():
    name = request.args.get('name', '').strip()
    if not name:
        return jsonify({"exists": False})
    conn = cursor = None
    try:
        conn   = config.get_db_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute(
            "SELECT file_id FROM uploaded_files WHERE stored_name = %s LIMIT 1",
            (name,)
        )
        row = cursor.fetchone()
        return jsonify({"exists": bool(row)})
    except Exception as e:
        print("check_excel_name error:", e)
        return jsonify({"exists": False})
    finally:
        if cursor: cursor.close()
        if conn:   conn.close()


# ── Rename uploaded file ──────────────────────────────────────────

@excel_bp.route('/rename_excel/<int:file_id>', methods=['POST'])
@login_required
def rename_excel(file_id):
    conn = cursor = None
    try:
        conn   = config.get_db_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT * FROM uploaded_files WHERE file_id = %s", (file_id,))
        record = cursor.fetchone()
        cursor.close(); conn.close()
        conn = cursor = None

        if not record:
            return jsonify({"success": False, "message": "File record not found."}), 404

        new_name = request.form.get('new_name', '').strip()
        if not new_name:
            return jsonify({"success": False, "message": "New name is required."}), 400
        if not new_name.lower().endswith('.xlsx'):
            new_name += '.xlsx'

        old_path = FolderAndFile.get_upload_path(record['stored_name'])
        new_path = FolderAndFile.get_upload_path(new_name)

        if os.path.exists(new_path) and new_name != record['stored_name']:
            return jsonify({"success": False, "message": "A file with that name already exists."}), 409

        if os.path.exists(old_path):
            os.rename(old_path, new_path)

        conn   = config.get_db_connection()
        cursor = conn.cursor()
        cursor.execute(
            "UPDATE uploaded_files SET stored_name = %s WHERE file_id = %s",
            (new_name, file_id)
        )
        conn.commit()

        LogHelper.log_action(
            "Rename Excel",
            f"User {current_user.id} renamed '{record['stored_name']}' to '{new_name}'"
        )
        return jsonify({"success": True, "message": "File renamed.", "new_name": new_name})

    except Exception as e:
        print("rename_excel error:", e)
        return jsonify({"success": False, "message": str(e)}), 500
    finally:
        if cursor: cursor.close()
        if conn:   conn.close()


# ── Delete uploaded file ──────────────────────────────────────────

@excel_bp.route('/delete_excel/<int:file_id>', methods=['POST'])
@login_required
def delete_excel(file_id):
    conn = cursor = None
    try:
        conn   = config.get_db_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT * FROM uploaded_files WHERE file_id = %s", (file_id,))
        record = cursor.fetchone()
        cursor.close(); conn.close()
        conn = cursor = None

        if not record:
            return jsonify({"success": False, "message": "File record not found."}), 404

        filepath = FolderAndFile.get_upload_path(record['stored_name'])
        if os.path.exists(filepath):
            try:
                _safe_remove(filepath)
            except PermissionError:
                return jsonify({
                    "success": False,
                    "message": "File is currently being processed. Please try again in a few seconds."
                }), 409
            except OSError as ose:
                return jsonify({
                    "success": False,
                    "message": f"Could not delete the file: {ose.strerror}"
                }), 500

        conn   = config.get_db_connection()
        cursor = conn.cursor()

        # Delete all data that was uploaded from this file (cascading order)
        # 1. payments linked to invoices from this file
        cursor.execute(
            "DELETE p FROM payment p "
            "JOIN invoices i ON p.Invoice_Id = i.Invoice_Id "
            "WHERE i.uploaded_file_id = %s", (file_id,))
        # 2. hold join rows linked to invoices from this file
        cursor.execute(
            "DELETE h FROM invoice_subcontractor_hold_join h "
            "JOIN invoices i ON h.Invoice_Id = i.Invoice_Id "
            "WHERE i.uploaded_file_id = %s", (file_id,))
        # 3. invoices, gst_releases, hold_release, credit_note tagged with this file
        cursor.execute("DELETE FROM invoices     WHERE uploaded_file_id = %s", (file_id,))
        cursor.execute("DELETE FROM gst_releases WHERE uploaded_file_id = %s", (file_id,))
        cursor.execute("DELETE FROM hold_release WHERE uploaded_file_id = %s", (file_id,))
        cursor.execute("DELETE FROM credit_note  WHERE uploaded_file_id = %s", (file_id,))
        # 4. finally remove the file record
        cursor.execute("DELETE FROM uploaded_files WHERE file_id = %s", (file_id,))
        conn.commit()

        LogHelper.log_action(
            "Delete Excel",
            f"User {current_user.id} deleted '{record['stored_name']}' and all associated data"
        )
        return jsonify({"success": True, "message": "File and all associated contractor data deleted."})

    except Exception as e:
        print("delete_excel error:", e)
        return jsonify({"success": False, "message": str(e)}), 500
    finally:
        if cursor: cursor.close()
        if conn:   conn.close()


# ── Excel History page ────────────────────────────────────────────

@excel_bp.route('/excel_history')
@login_required
def excel_history():
    uploads = _get_all_uploads()
    folder  = FolderAndFile.get_upload_folder()
    for f in uploads:
        f['status'] = 'Active' if os.path.exists(os.path.join(folder, f['stored_name'])) else 'File Missing'
    return render_template('excel_history.html', uploaded_files=uploads)


@excel_bp.route('/excel_data_sets')
@login_required
def data_sets():
    return redirect(url_for('excel.excel_history'))


# ── Replace a specific file (from Data Sets) ─────────────────────

@excel_bp.route('/replace_excel_data/<int:file_id>', methods=['POST'])
@login_required
def replace_excel_data(file_id):
    conn = cursor = None
    try:
        conn   = config.get_db_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT * FROM uploaded_files WHERE file_id = %s", (file_id,))
        record = cursor.fetchone()
        cursor.close(); conn.close()
        conn = cursor = None

        if not record:
            return jsonify({"success": False, "message": "File record not found."}), 404

        file = request.files.get('file')
        if not file or not (file.filename.endswith('.xlsx') or file.filename.endswith('.xls')):
            return jsonify({"success": False, "message": "Please upload a valid Excel file (.xlsx / .xls)."}), 400

        filepath = FolderAndFile.get_upload_path(record['stored_name'])
        try:
            _safe_save_upload(file, filepath)
        except PermissionError:
            return jsonify({
                "success": False,
                "message": "File is currently being processed. Please try again in a few seconds."
            }), 409
        finally:
            try:
                file.stream.close()
            except Exception:
                pass

        total_rows = _count_excel_rows(filepath)

        conn   = config.get_db_connection()
        cursor = conn.cursor()
        cursor.execute(
            "UPDATE uploaded_files SET updated_date = NOW(), total_rows = %s WHERE file_id = %s",
            (total_rows, file_id)
        )
        conn.commit()

        LogHelper.log_action(
            "Replace Excel Data",
            f"User {current_user.id} replaced '{record['stored_name']}'"
        )
        view_url = url_for('excel.show_table', filename=record['stored_name'],
                           file_id=file_id)
        return jsonify({
            "success":  True,
            "message":  f"'{record['stored_name']}' replaced successfully. Open it to preview and save.",
            "view_url": view_url,
        })

    except Exception as e:
        print("replace_excel_data error:", e)
        return jsonify({"success": False, "message": str(e)}), 500
    finally:
        if cursor: cursor.close()
        if conn:   conn.close()


# ── Upload a new version of an existing file (from Data Sets) ────

@excel_bp.route('/new_version_excel/<int:file_id>', methods=['POST'])
@login_required
def new_version_excel(file_id):
    conn = cursor = None
    try:
        conn   = config.get_db_connection()
        cursor = conn.cursor(dictionary=True)
        cursor.execute("SELECT * FROM uploaded_files WHERE file_id = %s", (file_id,))
        record = cursor.fetchone()
        cursor.close(); conn.close()
        conn = cursor = None

        if not record:
            return jsonify({"success": False, "message": "File record not found."}), 404

        file = request.files.get('file')
        if not file or not (file.filename.endswith('.xlsx') or file.filename.endswith('.xls')):
            return jsonify({"success": False, "message": "Please upload a valid Excel file (.xlsx / .xls)."}), 400

        new_stored_name = _auto_rename(record['stored_name'])
        filepath        = FolderAndFile.get_upload_path(new_stored_name)
        _safe_save_upload(file, filepath)
        try:
            file.stream.close()
        except Exception:
            pass
        total_rows = _count_excel_rows(filepath)

        new_file_id = _insert_upload(record['stored_name'], new_stored_name, total_rows)

        LogHelper.log_action(
            "Upload New Version",
            f"User {current_user.id} uploaded new version '{new_stored_name}' of '{record['stored_name']}'"
        )
        view_url = url_for('excel.show_table', filename=new_stored_name,
                           file_id=new_file_id or '')
        return jsonify({
            "success":  True,
            "message":  f"New version saved as '{new_stored_name}'. Open it to preview and save.",
            "new_name": new_stored_name,
            "view_url": view_url,
        })

    except Exception as e:
        print("new_version_excel error:", e)
        return jsonify({"success": False, "message": str(e)}), 500
    finally:
        if cursor: cursor.close()
        if conn:   conn.close()


# ═══════════════════════════════════════════════════════════════════
# BELOW: ORIGINAL ROUTES — LOGIC UNCHANGED
# ═══════════════════════════════════════════════════════════════════

# ── Show Excel Table ──────────────────────────────────────────────

@excel_bp.route('/show_table/<filename>')
def show_table(filename):
    global data

    # Resolve file_id from URL param or DB lookup
    file_id = request.args.get('file_id', type=int)
    if not file_id:
        rec = _find_by_stored_name(filename)
        file_id = rec['file_id'] if rec else None

    filepath = FolderAndFile.get_upload_path(filename)

    # Parse the Excel file
    parsed = _parse_excel_file(filepath)

    file_info        = parsed['file_info']
    variables        = parsed['variables']
    hold_columns_raw = parsed['hold_columns_raw']
    data             = parsed['data']

    errors             = []
    warnings           = []
    subcontractor_data = None
    state_data         = None
    district_data      = None
    block_data         = None
    hold_columns       = []
    hold_types_lookup  = {}

    connection = config.get_db_connection()

    if connection:
        try:
            cursor = connection.cursor(dictionary=True)

            # ---------------------------------------------------------
            # State
            # ---------------------------------------------------------
            state_data = _fuzzy_lookup_state(
                cursor,
                file_info['State']
            )

            if not state_data:
                errors.append(
                    f"State '{file_info['State']}' is not valid. Please add it."
                )

            # ---------------------------------------------------------
            # District
            # ---------------------------------------------------------
            if state_data:
                district_data = _fuzzy_lookup_district(
                    cursor,
                    file_info['District'],
                    state_data['State_Id']
                )

                if not district_data:
                    errors.append(
                        f"District '{file_info['District']}' is not valid "
                        f"under state '{file_info['State']}'."
                    )

            # ---------------------------------------------------------
            # Block
            # ---------------------------------------------------------
            if district_data:
                block_data, block_created = _fuzzy_lookup_block(
                    cursor,
                    file_info['Block'],
                    district_data['District_id'],
                    connection
                )

                if block_created:
                    warnings.append(
                        f"New block '{block_data['Block_Name']}' was "
                        f"automatically added under district "
                        f"'{district_data['District_Name']}'."
                    )

                elif not block_data:
                    errors.append(
                        f"Block '{file_info['Block']}' is not valid "
                        f"under district '{file_info['District']}'."
                    )

            # ---------------------------------------------------------
            # Subcontractor
            # ---------------------------------------------------------
            sub_name = (
                file_info.get('Subcontractor') or ''
            ).strip()

            if sub_name:
                cursor.callproc(
                    'GetSubcontractorByName',
                    [sub_name]
                )

                for result in cursor.stored_results():
                    subcontractor_data = result.fetchone()

                if not subcontractor_data:
                    cursor.callproc(
                        'SaveContractor',
                        [
                            sub_name,
                            None,
                            None,
                            None,
                            None,
                            None,
                            None,
                            None,
                            None
                        ]
                    )

                    connection.commit()

                    cursor.callproc(
                        'GetSubcontractorByName',
                        [sub_name]
                    )

                    for result in cursor.stored_results():
                        subcontractor_data = result.fetchone()

            else:
                errors.append(
                    "Contractor Name not detected in the Excel file. "
                    "Ensure the file has a 'Contractor Name:' metadata row "
                    "above the data table."
                )

            # =========================================================
            # GET ALL CONFIGURED HOLD TYPES
            # =========================================================

            cursor.callproc("GetAllHoldTypes")

            hold_types_data = []

            for ht in cursor.stored_results():
                hold_types_data = ht.fetchall()

            # Create lookup:
            #
            # SD_RR               -> 2
            # Testing_Deposit_RR  -> 3
            # Maintaninece_RR     -> 4
            # Extra_Hold          -> 1
            #
            hold_types_lookup = {
                _norm_key(row['hold_type']): row['hold_type_id']
                for row in hold_types_data
                if row.get('hold_type')
            }

            print(
                "Configured Hold Types:",
                hold_types_lookup
            )

            # =========================================================
            # IMPORTANT FIX
            #
            # _parse_excel_file() currently detects a hold column only
            # when the Excel column name contains "hold".
            #
            # Therefore:
            #
            # Extra_Hold          -> detected
            # SD_RR               -> NOT detected
            # Testing_Deposit_RR  -> NOT detected
            # Maintaninece_RR     -> NOT detected
            #
            # The following code additionally detects Excel columns
            # whose names match configured Hold Types in the database.
            # =========================================================

            detected_hold_columns = {
                _norm_key(hc['column_name'])
                for hc in hold_columns_raw
                if hc.get('column_name')
            }

            for column_name, column_number in variables.items():

                if not column_name:
                    continue

                normalized_column = _norm_key(column_name)

                # Already detected by:
                #
                # if 'hold' in raw.lower()
                #
                # Skip it so Extra_Hold is not added twice.
                if normalized_column in detected_hold_columns:
                    continue

                # -----------------------------------------------------
                # NEW FIX:
                # Match Excel column with configured Hold Type
                # -----------------------------------------------------
                if normalized_column in hold_types_lookup:

                    hold_columns_raw.append({
                        'column_name': column_name,
                        'column_number': column_number
                    })

                    detected_hold_columns.add(
                        normalized_column
                    )

                    print(
                        "Configured hold column detected: "
                        f"{column_name} -> "
                        f"hold_type_id="
                        f"{hold_types_lookup[normalized_column]}"
                    )

            # =========================================================
            # RESOLVE HOLD TYPE IDs
            # =========================================================

            _DPR_CANONICAL = 'dpr excess amount'

            hold_columns = []

            for hc in hold_columns_raw:

                col_name = hc['column_name']

                col_lower = _norm_key(col_name)

                hold_type_id = hold_types_lookup.get(
                    col_lower
                )
                                # -----------------------------------------------------
                # Existing DPR logic
                # -----------------------------------------------------
                if not hold_type_id and col_name and col_name.strip():

                    # If this column looks like a DPR hold,
                    # reuse canonical DPR hold type.
                    if 'dpr' in col_lower:

                        hold_type_id = hold_types_lookup.get(
                            _DPR_CANONICAL
                        )

                        if hold_type_id:

                            hold_types_lookup[
                                col_lower
                            ] = hold_type_id

                            print(
                                f"Hold col DPR-mapped: "
                                f"{col_name} -> "
                                f"hold_type_id={hold_type_id}"
                            )

                        else:

                            cursor.execute(
                                """
                                INSERT INTO hold_types
                                    (hold_type, include_in_hold)
                                VALUES (%s, %s)
                                """,
                                (
                                    'DPR Excess Amount',
                                    0
                                )
                            )

                            connection.commit()

                            hold_type_id = cursor.lastrowid

                            hold_types_lookup[
                                _DPR_CANONICAL
                            ] = hold_type_id

                            hold_types_lookup[
                                col_lower
                            ] = hold_type_id

                            warnings.append(
                                "New hold type 'DPR Excess Amount' "
                                "was automatically added."
                            )

                            print(
                                f"Hold col DPR-created: "
                                f"{col_name} -> "
                                f"hold_type_id={hold_type_id}"
                            )

                    # -------------------------------------------------
                    # Existing unknown hold type logic
                    # -------------------------------------------------
                    else:

                        args = [
                            col_name.strip(),
                            0
                        ]

                        cursor.callproc(
                            'InsertHoldType',
                            args
                        )

                        connection.commit()

                        cursor.execute(
                            """
                            SELECT hold_type_id
                            FROM hold_types
                            WHERE hold_type = %s
                            LIMIT 1
                            """,
                            (col_name.strip(),)
                        )

                        new_row = cursor.fetchone()

                        hold_type_id = (
                            new_row['hold_type_id']
                            if new_row
                            else None
                        )

                        if hold_type_id:

                            hold_types_lookup[
                                col_lower
                            ] = hold_type_id

                            warnings.append(
                                f"New hold type '{col_name}' "
                                f"was automatically added "
                                f"to the database."
                            )

                            print(
                                f"Hold col AUTO: "
                                f"{col_name} -> "
                                f"hold_type_id={hold_type_id}"
                            )

                        else:

                            print(
                                f"Hold col FAIL auto-create: "
                                f"{col_name}"
                            )

                else:

                    print(
                        f"Hold col OK: "
                        f"{col_name} -> "
                        f"hold_type_id={hold_type_id}"
                    )

                # -----------------------------------------------------
                # Add resolved hold column
                # -----------------------------------------------------
                hold_columns.append({
                    'column_name': col_name,
                    'column_number': hc['column_number'],
                    'hold_type_id': hold_type_id,
                })

            # =========================================================
            # DEBUG OUTPUT
            # =========================================================

            print("\n========== FINAL HOLD COLUMNS ==========")

            for hold in hold_columns:
                print(
                    f"Column: {hold['column_name']} | "
                    f"Column Number: {hold['column_number']} | "
                    f"Hold Type ID: {hold['hold_type_id']}"
                )

            print("========================================\n")

            cursor.close()

        except Exception as e:

            print(
                f"Database error: {e}"
            )

            return (
                f"Database operation failed: {e}",
                500
            )

        finally:

            connection.close()

    # =============================================================
    # Contractor validation
    # =============================================================

    metadata_contractor = (
        file_info.get('Subcontractor') or ''
    ).strip()

    if (
        'Contractor_Name' not in variables
        and not metadata_contractor
    ):
        errors.append(
            "Contractor Name not found in Excel columns or metadata. "
            "Add a 'Contractor Name:' row above the data table, "
            "or include a Contractor Name column."
        )

    # =============================================================
    # Upload history
    # =============================================================

    all_uploads = _get_all_uploads()

    # =============================================================
    # Render page
    # =============================================================

    return render_template(
        'show_excel_file.html',

        file_info=file_info,

        variables=variables,

        data=data,

        subcontractor_data=subcontractor_data,

        state_data=state_data,

        district_data=district_data,

        block_data=block_data,

        errors=errors,

        warnings=warnings,

        hold_columns=hold_columns,

        hold_counter=len(hold_columns),

        file_id=file_id,

        has_contractor=bool(subcontractor_data),

        all_uploads=all_uploads,

        current_filename=filename,
    )


# ── Save Data ─────────────────────────────────────────────────────

@excel_bp.route('/save_data', methods=['POST'])
def save_data():
    # Extract form data
    subcontractor_id = request.form.get("subcontractor_data")
    state_id   = request.form.get("state_data")
    district_id = request.form.get("district_data")
    block_id   = request.form.get("block_data")
    hold_columns = request.form.get("hold_columns")
    hold_counter = request.form.get("hold_counter")
    file_id = request.form.get("file_id", type=int)

    if not subcontractor_id:
        return jsonify({"error": "Contractor not identified. Please ensure the Excel file "
                                 "contains a valid 'Contractor Name' metadata row, "
                                 "then re-open the file and try again."}), 400

    # ── Always re-read the latest file from disk ───────────────────
    # This ensures Replace/New-Version/multi-tab usage all save current data,
    # not whatever was in memory when show_table() last ran.
    save_rows          = None
    variables_from_file = {}

    if file_id:
        _rec = None
        _conn = _cur = None
        try:
            _conn = config.get_db_connection()
            _cur  = _conn.cursor(dictionary=True)
            _cur.execute(
                "SELECT stored_name FROM uploaded_files WHERE file_id = %s LIMIT 1",
                (file_id,)
            )
            _rec = _cur.fetchone()
        except Exception as _e:
            print(f"save_data file lookup error: {_e}")
        finally:
            if _cur:  _cur.close()
            if _conn: _conn.close()

        if _rec:
            _filepath = FolderAndFile.get_upload_path(_rec['stored_name'])
            if os.path.exists(_filepath):
                try:
                    _parsed            = _parse_excel_file(_filepath)
                    save_rows          = _parsed['data']
                    variables_from_file = _parsed['variables']
                except Exception as _pe:
                    print(f"save_data re-read error: {_pe}")

    # Fall back to the in-memory global only if file re-read failed
    if save_rows is None:
        save_rows = data

    has_contractor_col = 'Contractor_Name' in variables_from_file

    if not save_rows:
        return jsonify({"error": "No data provided to save"}), 400

    if save_rows:
        connection = config.get_db_connection()
        cursor = connection.cursor()
        try:
            # ── Delete ALL existing data for this contractor, then re-insert fresh ──
            if subcontractor_id:
                cursor.execute(
                    "DELETE p FROM payment p "
                    "JOIN invoices i ON p.Invoice_Id = i.Invoice_Id "
                    "WHERE i.Subcontractor_Id = %s", (subcontractor_id,))
                cursor.execute(
                    "DELETE FROM invoices WHERE Subcontractor_Id = %s", (subcontractor_id,))
                cursor.execute(
                    "DELETE FROM gst_releases WHERE Contractor_ID = %s", (subcontractor_id,))
                cursor.execute(
                    "DELETE FROM hold_release WHERE Contractor_Id = %s", (subcontractor_id,))
                cursor.execute(
                    "DELETE FROM credit_note WHERE Contractor_Id = %s", (subcontractor_id,))
                cursor.execute(
                    "DELETE FROM invoice_subcontractor_hold_join WHERE Contractor_Id = %s", (subcontractor_id,))
                connection.commit()

            # ── Pre-capture max IDs so we know which rows we insert ──────
            def _max_id(table, pk):
                cursor.execute(f"SELECT COALESCE(MAX({pk}), 0) FROM `{table}`")
                return cursor.fetchone()[0]

            pre_inv  = _max_id('invoices',     'Invoice_Id')
            pre_gst  = _max_id('gst_releases', 'GST_Release_Id')
            pre_hr   = _max_id('hold_release', 'Hold_Release_Id')
            pre_cn   = _max_id('credit_note',  'Credit_Note_Id')

            saved_inv, saved_pay, skipped, duplicates, failed = 0, 0, 0, 0, 0
            failed_rows      = []
            total_rows_count = len(save_rows)
            _sort_order      = 0

            for entry in save_rows:
                # Skip total/summary/blank rows
                if _is_total_row(entry):
                    skipped += 1
                    continue
                _sort_order += 1

                # Validate Contractor_Name if the column exists in the Excel
                if has_contractor_col:
                    contractor_val = str(entry.get('Contractor_Name') or '').strip()
                    if not contractor_val:
                        row_num = entry.get('Row Number', '?')
                        failed_rows.append(f"Row {row_num} skipped: Contractor Name missing")
                        failed += 1
                        continue

                invoice_date = _parse_date(entry.get("Invoice_Date"))

                save_data = {
                    "PMC_No":              entry.get("PMC_No"),
                    "Invoice_Details":     entry.get("Invoice_Details") or entry.get("Invoice_Detail") or '',
                    "Work_Type":           'none',
                    "Invoice_Date":        invoice_date,
                    "Invoice_No":          _clean_invoice_no(entry.get("Invoice_No") or entry.get("Invoice_Number")),
                    "Basic_Amount":        _to_float(entry.get("Basic_Amount")),
                    "Debit_Amount":        _to_float(entry.get("Debit_Amount")),
                    "After_Debit_Amount":  _to_float(entry.get("After_Debit_Amount")),
                    "Amount":              _to_float(entry.get("Amount")),
                    "GST_Amount":          _to_float(entry.get("GST_Amount")),
                    "TDS_Amount":          _to_float(entry.get("TDS_Amount")),
                    "SD_Amount":           _to_float(entry.get("SD_Amount")),
                    "On_Commission":       _to_float(entry.get("On_Commission")),
                    "Hydro_Testing":       _to_float(entry.get("Hydro_Testing")),
                    "Hold_Amount":         _to_float(entry.get("Hold_Amount")),
                    "GST_SD_Amount":       _to_float(entry.get("GST_SD_Amount")),
                    "Final_Amount":        _to_float(entry.get("Final_Amount")),
                    "Payment_Amount":      _to_float(entry.get("Payment_Amount")),
                    "Total_Amount":        _to_float(entry.get("Total_Amount")),
                    "TDS_Payment_Amount":  _to_float(entry.get("TDS_Payment_Amount")),
                    "UTR":                 _clean_utr(entry.get("UTR") or entry.get("UTR_No")),
                }
                village_name, work_type = None, None
                village_id  = None
                PMC_No             = save_data.get('PMC_No')
                Invoice_Details    = save_data.get('Invoice_Details')
                Invoice_Date       = save_data.get('Invoice_Date')
                Invoice_No         = save_data.get('Invoice_No')
                Basic_Amount       = save_data.get('Basic_Amount')
                Debit_Amount       = save_data.get('Debit_Amount')
                After_Debit_Amount = save_data.get('After_Debit_Amount')
                Amount             = save_data.get('Amount')
                GST_Amount         = save_data.get('GST_Amount')
                TDS_Amount         = save_data.get('TDS_Amount')
                SD_Amount          = save_data.get('SD_Amount')
                On_Commission      = save_data.get('On_Commission')
                Hydro_Testing      = save_data.get('Hydro_Testing')
                GST_SD_Amount      = save_data.get('GST_SD_Amount')
                Final_Amount       = save_data.get('Final_Amount')
                Payment_Amount     = save_data.get('Payment_Amount')
                Total_Amount       = save_data.get('Total_Amount')
                TDS_Payment_Amount = save_data.get('TDS_Payment_Amount')
                UTR                = save_data.get('UTR')

                # Universal payment amount resolution:
                # Amount_with_TDS  = gross payment (before TDS withheld)
                # TDS_Payment_Amount = TDS withheld
                # Amount_after_TDS_Deduction = actual amount contractor receives (Total_Amount)
                #
                # Case 1: Excel has explicit Payment_Amount column (e.g. VP Enterprises)
                #   → TDS = Payment_Amount - Total_Amount
                # Case 2: No Payment_Amount column but TDS_Amount column exists (e.g. MS Construction)
                #   → Payment_Amount = Total_Amount + TDS_Amount, TDS = TDS_Amount
                # Case 3: No TDS anywhere → Amount_with_TDS = Amount_after_TDS_Deduction = Total_Amount
                _pa = float(Payment_Amount or 0)
                _ta = float(Total_Amount   or 0)
                _tds_inv = float(TDS_Amount        or 0)  # invoice-level TDS column
                _tds_pay = float(TDS_Payment_Amount or 0)  # dedicated payment TDS column

                if not _tds_pay:
                    if _pa > 0 and _ta > 0 and _pa > _ta:
                        # Explicit gross > net: TDS = difference
                        _tds_pay = round(_pa - _ta, 2)
                    elif _tds_inv > 0 and _ta > 0:
                        # No gross column: use invoice TDS as payment TDS
                        _tds_pay = _tds_inv
                        _pa = round(_ta + _tds_inv, 2)

                if not _pa and _ta:
                    _pa = round(_ta + _tds_pay, 2)

                Payment_Amount     = _pa
                TDS_Payment_Amount = _tds_pay

                # Check if Village_Name is a dedicated column in this Excel row
                _village_col = (entry.get('Village_Name') or entry.get('Village')
                                or entry.get('village_name') or entry.get('village') or '')
                if _village_col:
                    village_name = str(_village_col).strip().lower()

                _row_saved = False

                if Invoice_Details:
                    words = Invoice_Details.lower().split()
                    # Strip trailing punctuation (e.g. "Village-" → "village") before token matching
                    _clean_words = [w.strip(',-/;:.()[]') for w in words]
                    # Extract village name using fuzzy tokens to handle typos and attached punctuation
                    if not village_name:
                        _vil_word = next((w for w in _clean_words if w in _VILLAGE_TOKENS), None)
                        if _vil_word:
                            village_pos = _clean_words.index(_vil_word)
                            village_name = " ".join(_clean_words[:village_pos])
                    _has_village = village_name or any(w in _VILLAGE_TOKENS for w in _clean_words)
                    if 'work' in _clean_words:
                        work_pos = _clean_words.index('work')
                        if village_name:
                            _vil_w = next((w for w in _clean_words if w in _VILLAGE_TOKENS), None)
                            vp = _clean_words.index(_vil_w) if _vil_w else 0
                            work_type = " ".join(_clean_words[vp + 1:work_pos + 1])
                        else:
                            work_type = " ".join(_clean_words[:work_pos + 1])

                    normalized_details = re.sub(r'\s+', ' ', Invoice_Details.strip()).lower()

                    # ── Priority 1: Credit / Debit notes (must be checked before village+work
                    #    so rows like "Debit for DPR Work" are not mis-routed as invoices) ──
                    if any(kw in Invoice_Details.lower() for kw in ['credit', 'debit']):
                        cursor.callproc(
                            'AddCreditNoteFromExcel',
                            [
                                PMC_No, Invoice_Details, Basic_Amount, Debit_Amount, After_Debit_Amount,
                                GST_Amount, Amount, Final_Amount, Payment_Amount, Total_Amount, UTR,
                                subcontractor_id, Invoice_No
                            ]
                        )
                        _row_saved = True

                    # ── Priority 2: Full invoice with village + work pattern ────────────────
                    elif _has_village and 'work' in _clean_words:
                        if block_id and village_name:
                            cursor.callproc("GetVillageId", (block_id, village_name))
                            _vrow = None
                            for result in cursor.stored_results():
                                _vrow = result.fetchone()
                            village_id = _vrow[0] if _vrow else None
                            if not village_id:
                                cursor.callproc("SaveVillage", (village_name, block_id))
                                cursor.callproc("GetVillageId", (block_id, village_name))
                                _vrow = None
                                for result in cursor.stored_results():
                                    _vrow = result.fetchone()
                                village_id = _vrow[0] if _vrow else None

                        args = (
                            PMC_No, village_id, work_type, Invoice_Details, Invoice_Date, Invoice_No,
                            Basic_Amount, Debit_Amount, After_Debit_Amount, Amount, GST_Amount, TDS_Amount,
                            SD_Amount, On_Commission, Hydro_Testing, GST_SD_Amount, Final_Amount,
                            subcontractor_id, 0, _sort_order
                        )
                        results = cursor.callproc('SaveInvoice', args)
                        invoice_id = results[-1]
                        for _sr in cursor.stored_results():
                            _row = _sr.fetchone()
                            if _row and _row[0]:
                                invoice_id = _row[0]
                                break
                        if not invoice_id:
                            cursor.execute("SELECT LAST_INSERT_ID()")
                            _row = cursor.fetchone()
                            invoice_id = _row[0] if _row and _row[0] else None
                        saved_inv += 1

                        has_payment = UTR or Payment_Amount or TDS_Payment_Amount or Total_Amount
                        if has_payment:
                            # Guard NULL-UTR duplicates: skip if a null-UTR payment already exists for this invoice
                            if not UTR:
                                cursor.execute(
                                    "SELECT Payment_Id FROM payment WHERE Invoice_Id = %s AND UTR IS NULL LIMIT 1",
                                    (invoice_id,))
                                if cursor.fetchone():
                                    print(f"Null-UTR payment already exists for invoice {invoice_id}, skipping.")
                                    has_payment = False
                            if has_payment:
                                try:
                                    cursor.callproc(
                                        "SavePayment",
                                        (PMC_No, Invoice_No, Payment_Amount,
                                         TDS_Payment_Amount, Total_Amount, UTR, invoice_id)
                                    )
                                    saved_pay += 1
                                except Exception as pe:
                                    if '1062' in str(pe) or 'Duplicate' in str(pe):
                                        print(f"Duplicate UTR skipped: {UTR}")
                                    elif '1452' in str(pe) or 'foreign key' in str(pe).lower():
                                        print(f"FK error: invoice_id={invoice_id} not in invoice table, payment skipped for UTR={UTR}")
                                    else:
                                        raise

                        if isinstance(hold_columns, str):
                            try:
                                hold_columns = json.loads(hold_columns)
                            except Exception:
                                hold_columns = ast.literal_eval(hold_columns)
                        if isinstance(hold_columns, list) and all(isinstance(hold, dict) for hold in hold_columns):
                            for hold in hold_columns:
                                hold_column_name = hold.get('column_name')
                                hold_type_id = hold.get('hold_type_id')
                                if hold_column_name and hold_type_id:
                                    hold_amount = entry.get(hold_column_name)
                                    if hold_amount is not None and invoice_id:
                                        try:
                                            cursor.callproc('InsertHoldJoinData', [
                                                subcontractor_id, invoice_id,
                                                hold_type_id, hold_amount
                                            ])
                                            connection.commit()
                                            print(f"Hold join saved: inv={invoice_id} type={hold_type_id} amt={hold_amount}")
                                        except Exception as he:
                                            if '1452' in str(he) or 'foreign key' in str(he).lower():
                                                print(f"FK error: invoice_id={invoice_id} not in invoice table, hold join skipped")
                                            elif '1062' in str(he) or 'Duplicate' in str(he):
                                                print(f"Duplicate hold join skipped for invoice_id={invoice_id}")
                                            else:
                                                raise
                                    elif not invoice_id:
                                        print(f"Skipping hold join: invoice_id is invalid for row {entry.get('Row Number', '?')}")
                        else:
                            print("Hold columns data is not a valid list of dictionaries.")
                        _row_saved = True

                    # ── Hold release (only if not already saved as credit/debit or invoice) ─
                    if not _row_saved:
                        _hold_kws = [
                            'hold release note', 'hold release',
                            'excess hold', 'hold release amount', 'dpr excess hold amount',
                            'excess hold amount', 'multi to single layer bill', 'hold amount',
                            'logging report'
                        ]
                        if any(kw in normalized_details for kw in _hold_kws) or _is_hold_release_note(normalized_details):
                            try:
                                _hr_type_id, _hr_fixed = _resolve_hold_release_target(
                                    cursor, entry, subcontractor_id, PMC_No, Invoice_No, Final_Amount, Invoice_Details)
                                cursor.callproc(
                                    'AddHoldReleaseFromExcel',
                                    [PMC_No, Invoice_No, Invoice_Details, Basic_Amount, Final_Amount, Total_Amount, UTR,
                                     subcontractor_id, _hr_type_id, _hr_fixed])
                                connection.commit()
                            except Exception as hre:
                                if '1062' in str(hre) or 'Duplicate' in str(hre):
                                    print(f"Duplicate hold release UTR skipped: {UTR}")
                                else:
                                    raise
                            # Also save to invoices so the row appears in the output Excel
                            cursor.callproc("insertExtrapaymet", (PMC_No, Invoice_Details, Basic_Amount, Debit_Amount,
                                    After_Debit_Amount, GST_Amount, Amount, Final_Amount, subcontractor_id, _sort_order, Invoice_No))
                            invoice_id = None
                            for result in cursor.stored_results():
                                _r = result.fetchone()
                                invoice_id = _r[0] if _r else None
                            saved_inv += 1
                            if UTR or Payment_Amount or TDS_Payment_Amount or Total_Amount:
                                try:
                                    cursor.callproc("SavePayment",
                                        (PMC_No, Invoice_No, Payment_Amount, TDS_Payment_Amount, Total_Amount, UTR, invoice_id))
                                    saved_pay += 1
                                except Exception as pe:
                                    if '1062' in str(pe) or 'Duplicate' in str(pe):
                                        pass
                                    else:
                                        raise
                            _row_saved = True

                        # ── GST Release ───────────────────────────────────────────────────
                        elif any(kw in Invoice_Details.lower() for kw in ['gst', 'release', 'gst release note']) \
                                and 'hold' not in Invoice_Details.lower():
                            cursor.callproc(
                                'AddGSTReleaseFromExcel',
                                [PMC_No, Invoice_No, Basic_Amount, Final_Amount, Total_Amount, UTR, subcontractor_id, _sort_order, Invoice_Date]
                            )
                            connection.commit()
                            _row_saved = True

                # ── Catch-all: any row not handled above that has a PMC number ─────────
                # This captures rows whose Invoice_Details have keywords like 'work' but
                # no village match, or any other unrecognised description format.
                if not _row_saved and PMC_No:
                    _details_catchall = (Invoice_Details or "").strip()
                    cursor.callproc("insertExtrapaymet", (PMC_No, _details_catchall, Basic_Amount, Debit_Amount,
                            After_Debit_Amount, GST_Amount, Amount, Final_Amount, subcontractor_id, _sort_order, Invoice_No))
                    invoice_id = None
                    for result in cursor.stored_results():
                        row = result.fetchone()
                        invoice_id = row[0] if row else None
                    # insertExtrapaymet doesn't carry SD/OC/HT or VWA; patch them in if present
                    if invoice_id and (SD_Amount or On_Commission or Hydro_Testing):
                        cursor.execute(
                            "UPDATE invoices SET SD_Amount=%s, On_Commission=%s, Hydro_Testing=%s "
                            "WHERE Invoice_Id=%s",
                            (SD_Amount, On_Commission, Hydro_Testing, invoice_id)
                        )
                    saved_inv += 1
                    has_payment2 = UTR or Payment_Amount or TDS_Payment_Amount or Total_Amount
                    if has_payment2:
                        try:
                            cursor.callproc(
                                "SavePayment",
                                (PMC_No, Invoice_No, Payment_Amount,
                                 TDS_Payment_Amount, Total_Amount, UTR, invoice_id)
                            )
                            saved_pay += 1
                        except Exception as pe:
                            if '1062' in str(pe) or 'Duplicate' in str(pe):
                                print(f"Duplicate UTR skipped: {UTR}")
                            elif '1452' in str(pe) or 'foreign key' in str(pe).lower():
                                print(f"FK error: invoice_id={invoice_id} not in invoice table, payment skipped for UTR={UTR}")
                            else:
                                raise
                    # Save hold join data for catch-all invoice rows
                    if isinstance(hold_columns, str):
                        try:
                            hold_columns = json.loads(hold_columns)
                        except Exception:
                            hold_columns = ast.literal_eval(hold_columns)
                    if isinstance(hold_columns, list) and all(isinstance(hc, dict) for hc in hold_columns):
                        for hc in hold_columns:
                            hc_name    = hc.get('column_name')
                            hc_type_id = hc.get('hold_type_id')
                            if hc_name and hc_type_id:
                                hc_amount = entry.get(hc_name)
                                if hc_amount is not None and invoice_id:
                                    try:
                                        cursor.callproc('InsertHoldJoinData', [
                                            subcontractor_id, invoice_id, hc_type_id, hc_amount
                                        ])
                                        connection.commit()
                                    except Exception as hce:
                                        if '1452' in str(hce) or 'foreign key' in str(hce).lower():
                                            print(f"FK error: hold join skipped for invoice_id={invoice_id}")
                                        elif '1062' in str(hce) or 'Duplicate' in str(hce):
                                            print(f"Duplicate hold join skipped for invoice_id={invoice_id}")
                                        else:
                                            raise

            # ── Tag all newly inserted rows with this file_id ────────────
            if file_id:
                def _tag(table, pk, pre):
                    cursor.execute(
                        f"UPDATE `{table}` SET uploaded_file_id=%s "
                        f"WHERE {pk} > %s AND uploaded_file_id IS NULL",
                        (file_id, pre))

                _tag('invoices',     'Invoice_Id',     pre_inv)
                _tag('gst_releases', 'GST_Release_Id', pre_gst)
                _tag('hold_release', 'Hold_Release_Id',pre_hr)
                _tag('credit_note',  'Credit_Note_Id', pre_cn)

            connection.commit()
            LogHelper.log_action("Data saved",
                f"User {current_user.id}: {saved_inv} invoices, {saved_pay} payments saved, "
                f"{duplicates} duplicates skipped, {skipped} blank rows skipped, {failed} failed rows.")
            msg = (f"Data saved successfully! "
                   f"{saved_inv} invoice record{'s' if saved_inv!=1 else ''}, "
                   f"{saved_pay} payment record{'s' if saved_pay!=1 else ''} saved.")
            if duplicates:
                msg += f" {duplicates} duplicate row{'s' if duplicates!=1 else ''} skipped."
            if skipped:
                msg += f" {skipped} blank/total row{'s' if skipped!=1 else ''} skipped."
            if failed:
                msg += f" {failed} row{'s' if failed!=1 else ''} failed validation."
            return jsonify({"success": msg, "summary": {
                "total_rows": total_rows_count,
                "saved_invoices": saved_inv,
                "saved_payments": saved_pay,
                "duplicates_skipped": duplicates,
                "blank_rows_skipped": skipped,
                "failed_rows": failed,
                "failed_details": failed_rows,
            }}), 200
        except Exception as e:
            connection.rollback()
            return jsonify({"error": f"An unexpected error occurred: {e}"}), 500
        finally:
            cursor.close()
            connection.close()
    return render_template('index.html')