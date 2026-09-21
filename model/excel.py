import openpyxl
from openpyxl.styles import Font, PatternFill
from datetime import datetime
from decimal import Decimal
from model.FolderAndFile import FolderAndFile

class excel:

    @staticmethod
    def generate_excel(contractor_id, info, rows, hold_types, hold_data, credit_note_map, payments, total, output_file):
        """
        rows  – unified list of invoice + GST release dicts, already sorted in
                source Excel row order (sort_order column).  Each dict has
                _row_type = 'invoice' | 'gst_release'.
        """
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Contractor Report"

        # ── Contractor Info ────────────────────────────────────────────────
        for field, value in info.items():
            sheet.append([field.replace("_", " "), value])
        sheet.append([])

        # ── Headers ───────────────────────────────────────────────────────
        base_headers = [
            "PMC No", "Village", "Work Type", "Invoice Details",
            "Invoice Date", "Invoice No",
            "Basic Amount", "Debit", "After Debit Amount", "GST (18%)",
            "Amount", "TDS (1%)", "SD (5%)", "On Commission",
            "Hydro Testing", "GST SD Amount", "Village Wise Advance"
        ]
        hold_headers = [ht['hold_type'] for ht in hold_types]
        payment_headers = ["Final Amount", "Payment Amount", "TDS Payment", "Total Paid", "UTR"]
        all_headers = base_headers + hold_headers + payment_headers

        sheet.append(all_headers)
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid")

        # ── Village Wise Advance per PMC group ───────────────────────────
        # Formula: sum(Final_Amount) - sum(Total_Amount) for all rows in the PMC
        # Value shown only on the LAST row of each PMC group.
        from collections import Counter, defaultdict
        _pmc_final  = defaultdict(float)
        _pmc_paid   = defaultdict(float)
        _pmc_last_idx = {}
        for _i, _r in enumerate(rows):
            _p = str(_r.get("PMC_No") or "").strip()
            _pmc_final[_p] += float(_r.get("Final_Amount") or 0)
            _pmc_paid[_p]  += float(_r.get("Total_Amount") or 0)
            _pmc_last_idx[_p] = _i
        _vwa_by_pmc = {p: round(_pmc_final[p] - _pmc_paid[p], 2) for p in _pmc_final}

        # ── Detect duplicate UTRs ─────────────────────────────────────────
        _all_utrs = []
        for _r in rows:
            _u = str(_r.get("UTR") or "").strip()
            if _u: _all_utrs.append(_u)
        _dup_utrs = {u for u, c in Counter(_all_utrs).items() if c > 1}
        _dup_fill = PatternFill(start_color="FF6600", end_color="FF6600", fill_type="solid")
        # UTR is the last column: 17 base + N hold + 5 payment cols
        _utr_col  = 17 + len(hold_headers) + 5

        # ── Row loop ──────────────────────────────────────────────────────
        appended_credit_keys = set()
        previous_pmc_no = None

        for idx, row in enumerate(rows):
            pmc_no   = str(row.get("PMC_No") or "").strip()
            row_type = row.get("_row_type", "invoice")

            # ── Yellow separator + credit notes when PMC section ends ──────
            if previous_pmc_no and pmc_no != previous_pmc_no:
                # Flush credit notes for the PMC that just ended
                if previous_pmc_no in credit_note_map and previous_pmc_no not in appended_credit_keys:
                    for cn in credit_note_map[previous_pmc_no]:
                        cn_row = [
                            previous_pmc_no, "", "",
                            cn.get("Invoice_Details", "Credit Note"), "",
                            cn.get("Invoice_No", ""),
                            cn.get("Basic_Amount", ""),
                            cn.get("Debit_Amount", ""),
                            cn.get("After_Debit_Amount", ""),
                            cn.get("GST_Amount", ""),
                            cn.get("Amount", ""),
                            "", "", "", "", "", ""
                        ]
                        cn_row += [""] * len(hold_headers)
                        cn_row += [cn.get("Final_Amount", ""), "", "", cn.get("Total_Amount", ""), cn.get("UTR", "")]
                        sheet.append(cn_row)
                    appended_credit_keys.add(previous_pmc_no)

                # Yellow separator row
                sheet.append([""] * len(all_headers))
                yellow_fill = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")
                for cell in sheet[sheet.max_row]:
                    cell.fill = yellow_fill

            previous_pmc_no = pmc_no

            # ── Invoice row ───────────────────────────────────────────────
            if row_type == "invoice":
                raw_inv_no = row.get("invoice_no")
                invoice_no = (
                    str(raw_inv_no).replace(" ", "")
                    if raw_inv_no not in (None, "", 0)
                    else ""
                )

                data_row = [
                    pmc_no,
                    row.get("Village_Name", ""),
                    row.get("Work_Type", ""),
                    row.get("Invoice_Details", ""),
                    row.get("Invoice_Date", ""),
                    invoice_no,
                    row.get("Basic_Amount", ""),
                    row.get("Debit_Amount", ""),
                    row.get("After_Debit_Amount", ""),
                    row.get("GST_Amount", ""),
                    row.get("Amount", ""),
                    row.get("TDS_Amount", ""),
                    row.get("SD_Amount", ""),
                    row.get("On_Commission", ""),
                    row.get("Hydro_Testing", ""),
                    row.get("GST_SD_Amount", ""),
                    _vwa_by_pmc.get(pmc_no, "") if idx == _pmc_last_idx.get(pmc_no) else "",
                ]

                invoice_holds = hold_data.get(row.get("Invoice_Id"), {})
                for ht in hold_types:
                    data_row.append(invoice_holds.get(ht['hold_type_id'], ""))

                data_row += [
                    row.get("Final_Amount", ""),
                    row.get("Payment_Amount", ""),
                    row.get("TDS_Payment_Amount", ""),
                    row.get("Total_Amount", ""),
                    row.get("UTR", ""),
                ]
                sheet.append(data_row)

                # Highlight UTR cell orange if duplicate
                _utr_val = str(row.get("UTR") or "").strip()
                if _utr_val and _utr_val in _dup_utrs:
                    _c = sheet.cell(row=sheet.max_row, column=_utr_col)
                    _c.fill = _dup_fill
                    _c.font = Font(bold=True, color="FFFFFF")

            # ── GST Release row ───────────────────────────────────────────
            elif row_type == "gst_release":
                inv_no = row.get("invoice_no") or row.get("Invoice_No", "")
                gst_row = [
                    pmc_no, "", "", "GST Release Note", "", inv_no,
                    row.get("Basic_Amount", ""),
                    "", "", "", "", "", "", "", "",
                    "", _vwa_by_pmc.get(pmc_no, "") if idx == _pmc_last_idx.get(pmc_no) else ""
                ]
                gst_row += [""] * len(hold_headers)
                gst_row += [
                    row.get("Final_Amount", ""),
                    "", "",
                    row.get("Total_Amount", ""),
                    row.get("UTR", ""),
                ]
                sheet.append(gst_row)

        # ── Flush credit notes for the last PMC ───────────────────────────
        if previous_pmc_no and previous_pmc_no in credit_note_map and previous_pmc_no not in appended_credit_keys:
            for cn in credit_note_map[previous_pmc_no]:
                cn_row = [
                    previous_pmc_no, "", "",
                    cn.get("Invoice_Details", "Credit Note"), "",
                    cn.get("Invoice_No", ""),
                    cn.get("Basic_Amount", ""),
                    cn.get("Debit_Amount", ""),
                    cn.get("After_Debit_Amount", ""),
                    cn.get("GST_Amount", ""),
                    cn.get("Amount", ""),
                    "", "", "", "", "", ""
                ]
                cn_row += [""] * len(hold_headers)
                cn_row += [cn.get("Final_Amount", ""), "", "", cn.get("Total_Amount", ""), cn.get("UTR", "")]
                sheet.append(cn_row)
            appended_credit_keys.add(previous_pmc_no)

        # ── Total row ─────────────────────────────────────────────────────
        hold_totals = {ht['hold_type_id']: 0 for ht in hold_types}
        for row in rows:
            if row.get('_row_type') == 'invoice':
                invoice_holds = hold_data.get(row.get("Invoice_Id"), {})
                for ht_id in hold_totals:
                    hold_totals[ht_id] += invoice_holds.get(ht_id, 0) or 0

        total_basic_amount      = total.get('sum_invo_basic_amt', 0) + total.get('sum_gst_basic_amt', 0) + total.get('sum_credit_basic_amt', 0)
        total_final_amount      = total.get('sum_invo_final_amt', 0) + total.get('sum_gst_final_amt', 0) + total.get('sum_credit_final_amt', 0)
        total_total_amount      = total.get('sum_pay_total_amt', 0)  + total.get('sum_gst_total_amt', 0) + total.get('sum_credit_total_amt', 0)
        total_gst_amount        = total.get('sum_invo_gst_amt', 0)   + total.get('sum_credit_gst_amt', 0)
        total_debit_amount      = total.get('sum_invo_debit_amt', 0) + total.get('sum_credit_debit_amount_amt', 0)
        total_after_debit_amount = total.get('sum_invo_after_debit_amt', 0) + total.get('sum_credit_after_debit_amt', 0)

        totalrow = [
            "Total", "-", "-", "-", "-", "-",
            total_basic_amount,
            total_debit_amount,
            total_after_debit_amount,
            total_gst_amount,
            total.get('sum_invo_amt', 0),
            total.get('sum_invo_tds_amt', 0),
            total.get('sum_invo_ds_amt', 0),
            total.get('sum_invo_on_commission', 0),
            total.get('sum_invo_hydro_test', 0),
            total.get('sum_invo_gst_sd_amt', 0),
            "",  # Village Wise Advance — not summed
        ]
        for ht in hold_types:
            totalrow.append(hold_totals.get(ht['hold_type_id'], 0))
        totalrow += [total_final_amount, "", "", total_total_amount]

        sheet.append(totalrow)
        for cell in sheet[sheet.max_row]:
            cell.fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
            cell.font = Font(bold=True)

        # ── Summary section ───────────────────────────────────────────────
        today_date = datetime.today().strftime('%A, %Y-%m-%d')
        sheet.append([])
        sheet.append(["Summary Report"])
        sheet.append([])
        sheet.append(["Contractor Name", info["Contractor_Name"]])
        sheet.append(["Date", today_date])
        sheet.append(["Description", "Amount"])

        # ── Summary rows ──────────────────────────────────────────────────
        sd_amt = float(total.get('sum_invo_ds_amt', 0) or 0)
        oc_amt = float(total.get('sum_invo_on_commission', 0) or 0)
        ht_amt = float(total.get('sum_invo_hydro_test', 0) or 0)

        # Hold = SD + OC + HT + non-DPR custom holds (include_in_hold=1), matching source Excel formula
        custom_total = sum(
            float(hold_totals.get(ht['hold_type_id'], 0) or 0)
            for ht in hold_types if ht.get('include_in_hold', 1)
        )
        hold_amt = round(sd_amt + oc_amt + ht_amt + custom_total, 2)
        sheet.append(["Hold", hold_amt])
        hold_row = sheet.max_row
        for col_idx in (1, 2):
            c = sheet.cell(row=hold_row, column=col_idx)
            c.font = Font(bold=True)
            c.fill = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")

        # Total Debit
        total_debit = round(float(total.get('sum_invo_debit_amt', 0) or 0), 2)
        sheet.append(["Total Debit", total_debit])

        # Advance / Surplus
        surplus_amount = round(total_final_amount - total_total_amount, 2)
        sheet.append(["Advance/ Surplus", surplus_amount])
        current_row = sheet.max_row
        cell = sheet.cell(row=current_row, column=2)
        cell.fill = PatternFill(
            start_color="C6EFCE" if surplus_amount >= 0 else "FFC7CE",
            end_color="C6EFCE"   if surplus_amount >= 0 else "FFC7CE",
            fill_type="solid"
        )

        # GST Remaining = invoice GST_Amount total - GST already raised as release notes
        gst_remaining = round(
            float(total.get('sum_invo_gst_amt', 0) or 0) -
            float(total.get('sum_gst_final_amt',  0) or 0), 2)
        sheet.append(["GST Remaining", gst_remaining])

        # Custom hold types shown as gross — hold releases are already captured in payments/surplus
        for ht in hold_types:
            gross = round(float(hold_totals.get(ht['hold_type_id'], 0) or 0), 2)
            if gross:
                sheet.append([ht['hold_type'], gross])

        # ── Auto-fit columns ──────────────────────────────────────────────
        for col in sheet.columns:
            max_length = 0
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            for cell in col:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            sheet.column_dimensions[col_letter].width = max_length + 2

        workbook.save(output_file)
